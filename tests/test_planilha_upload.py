"""Testes de upload/parsing de planilha (PUT /planilha/upload)."""
import io
import uuid

from openpyxl import Workbook

from database.models import HistoricDashboard, PlanilhaModel

XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument"
    ".spreadsheetml.sheet"
)

HEADERS = [
    "NOME DO PRODUTO",
    "DATA DE VENDA (DIA-MÊS-ANO)",
    "DATA DO PAGAMENTO (DIA-MÊS-ANO)",
    "VALOR BRUTO (R$)",
    "VALOR LIQUIDO (R$)",
    "TAXA (R$)",
    "FORMA DE PAGAMENTO",
    "CATEGORIA",
]


def _username():
    return "u" + uuid.uuid4().hex[:12]


def _build_planilha_xlsx(rows):
    """Monta um xlsx no mesmo layout esperado por create_planilha (11 linhas
    de cabeçalho decorativo puladas via skiprows=range(0, 11), header na
    linha 12)."""
    wb = Workbook()
    ws = wb.active
    for _ in range(11):
        ws.append([])
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def test_upload_valid_planilha_inserts_rows(client, make_user, auth_token, db_session):
    user, _ = make_user(_username())
    token = auth_token(user.username)

    xlsx = _build_planilha_xlsx([
        ["Produto 1", "10-01-2026", "15-01-2026", 100.0, 90.0, 10.0, "Pix", "Categoria A"],
        ["Produto 2", "11-01-2026", "16-01-2026", 200.0, 180.0, 20.0, "Crédito", "Categoria B"],
    ])

    resp = client.put(
        "/planilha/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"selected_file": ("planilha.xlsx", xlsx, XLSX_MIME)},
    )
    assert resp.status_code == 200

    rows = db_session.query(PlanilhaModel).filter(
        PlanilhaModel.user_id == user.id
    ).all()
    assert len(rows) == 2
    assert {r.nome_produto for r in rows} == {"Produto 1", "Produto 2"}

    historico = db_session.query(HistoricDashboard).filter(
        HistoricDashboard.user_id == user.id
    ).all()
    assert len(historico) == 1
    assert all(r.historic_dashboard_id == historico[0].id for r in rows)


def test_upload_with_missing_columns_returns_400_not_500(client, make_user, auth_token):
    """Planilha sem as colunas esperadas não deve derrubar o servidor com um
    erro não tratado (ver achado de Erros: KeyError do pandas cai no
    `except Exception` genérico e vira 400, mas sem log do motivo real)."""
    user, _ = make_user(_username())
    token = auth_token(user.username)

    wb = Workbook()
    ws = wb.active
    for _ in range(11):
        ws.append([])
    ws.append(["coluna_errada_1", "coluna_errada_2"])
    ws.append(["valor 1", "valor 2"])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    resp = client.put(
        "/planilha/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"selected_file": ("planilha.xlsx", buffer, XLSX_MIME)},
    )
    assert resp.status_code == 400
    # não deve vazar a mensagem crua da exceção do pandas/SQLAlchemy
    assert "Traceback" not in resp.text
    assert "KeyError" not in resp.text


def test_upload_does_not_create_orphan_historic_when_parsing_fails(
    client, make_user, auth_token, db_session
):
    """Antes desta correção, o HistoricDashboard era commitado ANTES do
    parsing, então uma planilha inválida ainda deixava um registro de
    histórico "vazio" no banco. Agora o histórico só é criado depois que o
    parsing dá certo."""
    user, _ = make_user(_username())
    token = auth_token(user.username)

    wb = Workbook()
    ws = wb.active
    for _ in range(11):
        ws.append([])
    ws.append(["coluna_errada"])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    resp = client.put(
        "/planilha/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"selected_file": ("planilha.xlsx", buffer, XLSX_MIME)},
    )
    assert resp.status_code == 400

    historico = db_session.query(HistoricDashboard).filter(
        HistoricDashboard.user_id == user.id
    ).all()
    assert historico == []
